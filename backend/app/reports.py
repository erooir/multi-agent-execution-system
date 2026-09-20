"""Versioned reports and safe local exports."""

from __future__ import annotations

import html
import io
import re
import threading
from copy import deepcopy
from uuid import uuid4

from .storage import store

_lock = threading.RLock()
REPORT_TEMPLATES = [
    {
        "id": "technology",
        "name": "科技情报专题报告",
        "sections": ["任务范围", "资料与方法", "研究发现", "证据限制", "人工复核", "参考资料"],
    },
    {
        "id": "geography",
        "name": "地理情报分析报告",
        "sections": ["研究范围", "数据来源", "跨期对比", "数据质量", "人工复核", "参考资料"],
    },
    {
        "id": "situational",
        "name": "态势研判报告",
        "sections": ["事件摘要", "时间线", "来源核验", "不确定事项", "人工复核", "参考资料"],
    },
    {
        "id": "summary",
        "name": "任务执行摘要报告",
        "sections": ["任务目的", "执行过程", "主要结果", "异常与待办", "人工复核", "参考资料"],
    },
]


def _find(report_id: str) -> dict:
    report = store.get("reports", report_id)
    if not report:
        raise ValueError("报告不存在")
    return report


def _snapshot(report: dict) -> dict:
    identity = f"{report['id']}:{report['version']}"
    existing = store.get("report_versions", identity)
    if existing:
        return existing
    return store.save(
        "report_versions",
        {
            "id": identity,
            "report_id": report["id"],
            "version": report["version"],
            "snapshot": deepcopy(report),
            "content": report["content"],
            "title": report["title"],
            "status": report["status"],
        },
    )


def _validate_content(content: str) -> str:
    if not isinstance(content, str) or not content.strip():
        raise ValueError("报告内容不能为空")
    if len(content) > 500_000:
        raise ValueError("报告内容超过 50 万字符限制")
    return content.strip()


def _citations(evidence: list[dict]) -> list[dict]:
    citations = []
    seen = set()
    for item in evidence:
        chunk_id = item.get("id")
        if not chunk_id or chunk_id in seen:
            continue
        chunk = store.get("chunks", chunk_id)
        if not chunk or chunk.get("document_id") != item.get("document_id"):
            raise ValueError("报告包含无法定位的引用，拒绝保存伪造证据")
        document = store.get("documents", chunk["document_id"])
        if not document:
            raise ValueError("引用资料不存在")
        seen.add(chunk_id)
        citations.append(
            {
                "id": chunk_id,
                "document_id": chunk["document_id"],
                "document_name": document["name"],
                "location": chunk["location"],
                "text": chunk["text"],
                "visibility": document["visibility"],
                "synthetic": document.get("synthetic", False),
            }
        )
    return citations


def create_report(run: dict, content: str, evidence: list[dict]) -> dict:
    content = _validate_content(content)
    citations = _citations(evidence)
    rehearsal = run.get("mode") == "rehearsal"
    disclosure = "> 本报告来自离线演练：使用真实本地检索与确定性格式整理，未调用外部模型。"
    if rehearsal and "本报告来自离线演练" not in content:
        content = disclosure + "\n\n" + content
    if any(item.get("synthetic") for item in citations) and "合成" not in content[:400]:
        content = "> 资料声明：报告引用合成演示资料，不代表真实研究结论、事件或地理信息。\n\n" + content
    project = store.get("projects", run.get("project_id", "")) or {}
    with _lock:
        report = store.save(
            "reports",
            {
                "id": "report_" + uuid4().hex,
                "title": run.get("report_title") or run.get("name") or "研究报告",
                "category": run.get("category") or project.get("category", "technology"),
                "project_id": run.get("project_id"),
                "run_id": run.get("id"),
                "content": content,
                "citations": citations,
                "version": 1,
                "status": "reviewed" if run.get("reviewed_content") else "draft",
                "mode": run.get("mode", "rehearsal"),
                "template": run.get("report_template", project.get("category", "technology")),
            },
        )
        _snapshot(report)
        store.audit(
            "report.create",
            report["id"],
            {"run_id": run.get("id"), "version": 1, "citations": len(citations)},
        )
        return report


def update_report(report_id: str, content: str, title: str | None = None) -> dict:
    content = _validate_content(content)
    if title is not None and (not isinstance(title, str) or not title.strip() or len(title) > 200):
        raise ValueError("报告标题应为 1 至 200 个字符")
    with _lock:
        previous = _find(report_id)
        _snapshot(previous)
        report = store.save(
            "reports",
            {
                **previous,
                "content": content,
                "title": title.strip() if title else previous["title"],
                "version": previous["version"] + 1,
                "status": "draft",
            },
        )
        _snapshot(report)
        store.audit(
            "report.update",
            report_id,
            {"version": report["version"], "previous_version": previous["version"]},
        )
        return report


def versions(report_id: str) -> list[dict]:
    _find(report_id)
    return sorted(
        [version for version in store.list("report_versions") if version.get("report_id") == report_id],
        key=lambda version: version["version"],
        reverse=True,
    )


def mark_reviewed(report_id: str, content: str | None = None, feedback: str | None = None) -> dict:
    """Approval is a new immutable version, including an optional reviewer edit."""
    with _lock:
        previous = _find(report_id)
        _snapshot(previous)
        reviewed = store.save(
            "reports",
            {
                **previous,
                "content": _validate_content(content) if content is not None else previous["content"],
                "status": "reviewed",
                "review_feedback": feedback or "",
                "version": previous["version"] + 1,
            },
        )
        _snapshot(reviewed)
        store.audit("report.approve", report_id, {"version": reviewed["version"], "feedback": feedback or ""})
        return reviewed


def restore(report_id: str, version: int) -> dict:
    with _lock:
        previous = _find(report_id)
        source = store.get("report_versions", f"{report_id}:{int(version)}")
        if not source:
            raise ValueError("报告历史版本不存在")
        _snapshot(previous)
        snapshot = source["snapshot"]
        restored = store.save(
            "reports",
            {
                **previous,
                "title": snapshot["title"],
                "content": snapshot["content"],
                "citations": deepcopy(snapshot.get("citations", [])),
                "status": "draft",
                "version": previous["version"] + 1,
                "restored_from": int(version),
            },
        )
        _snapshot(restored)
        store.audit(
            "report.restore", report_id, {"source_version": int(version), "new_version": restored["version"]}
        )
        return restored


def _with_sources(report: dict) -> str:
    content = report["content"]
    if not content.lstrip().startswith("# "):
        content = "# " + report["title"] + "\n\n" + content
    metadata = f"报告版本：v{report['version']} · 状态：{report['status']} · 模式：{'离线演练' if report.get('mode') == 'rehearsal' else '真实模型运行'}"
    content += "\n\n---\n\n" + metadata
    if report.get("citations"):
        content += "\n\n## 引用资料索引\n"
        for index, citation in enumerate(report["citations"], 1):
            content += f"\n{index}. [{citation['id']}] {citation['document_name']} — {citation['location']}\n"
    return content + "\n"


def _safe_html(markdown: str) -> str:
    blocks = []
    in_code = False
    for line in markdown.splitlines():
        if line.startswith("```"):
            blocks.append("</pre>" if in_code else "<pre>")
            in_code = not in_code
            continue
        escaped = html.escape(line)
        if in_code:
            blocks.append(escaped + "\n")
        elif line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 6)
            blocks.append(f"<h{level}>{html.escape(line[level:].strip())}</h{level}>")
        elif line.startswith("> "):
            blocks.append(f"<blockquote>{html.escape(line[2:])}</blockquote>")
        elif line.strip() == "---":
            blocks.append("<hr>")
        elif line:
            escaped = re.sub(r"\*\*([^*]+)\*\*", r"<strong>\1</strong>", escaped)
            blocks.append("<p>" + escaped + "</p>")
    if in_code:
        blocks.append("</pre>")
    return "\n".join(blocks)


def export_report(report_id: str, format: str) -> tuple[bytes, str, str]:
    report = _find(report_id)
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", report["title"]).strip(" .")[:100] or "研究报告"
    markdown = _with_sources(report)
    if format == "md":
        return markdown.encode("utf-8"), "text/markdown; charset=utf-8", f"{name}_v{report['version']}.md"
    if format == "html":
        page = "<!doctype html><html lang='zh-CN'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        page += (
            "<title>"
            + html.escape(report["title"])
            + "</title><style>body{max-width:900px;margin:40px auto;padding:0 24px;color:#142537;font:16px/1.85 system-ui,sans-serif}h1,h2,h3{line-height:1.4}blockquote{padding:12px 18px;background:#eef5fa;border-left:3px solid #45748c}pre{white-space:pre-wrap;overflow-wrap:anywhere}p{overflow-wrap:anywhere}</style></head><body>"
        )
        page += _safe_html(markdown) + "</body></html>"
        return page.encode("utf-8"), "text/html; charset=utf-8", f"{name}_v{report['version']}.html"
    if format == "docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor

        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.8)
        style = document.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(7)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        for line in markdown.splitlines():
            if not line.strip() or line.strip() == "---":
                continue
            if line.startswith("# "):
                paragraph = document.add_heading(line[2:].strip(), 0)
            elif re.match(r"^#{2,6} ", line):
                level = min(len(line) - len(line.lstrip("#")) - 1, 3)
                paragraph = document.add_heading(line.lstrip("#").strip(), level)
            elif line.startswith(("- ", "* ")):
                paragraph = document.add_paragraph(line[2:], "List Bullet")
            else:
                paragraph = document.add_paragraph(line.removeprefix("> ").replace("**", ""))
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor.from_string("142537")
                fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
                fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        output = io.BytesIO()
        document.save(output)
        return (
            output.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            f"{name}_v{report['version']}.docx",
        )
    raise ValueError("导出格式应为 md、docx 或 html")
